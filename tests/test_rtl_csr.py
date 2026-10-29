"""The `csr` step of the RTL flow: the APB-CSR-Generator driven from a register map drafted from the spec (docs/spec/rtl-flow.md "CSR").
A fake generator (reads the real workbook with openpyxl, writes a small module) stands in for the tool; the real one is tried in
`test_the_real_generator` when it is installed ($CSR_GEN_HOME or ~/.q3tui/tools/APB-CSR-Generator)."""

import json
import sys
import time
from pathlib import Path

import anyio
import openpyxl
import pytest

from q3tui.core.project import Project
from q3tui.flows.rtl.lib import csr
from q3tui.flows.rtl.lib.models import CsrMap
from q3tui.llm import runtime
from q3tui.pipeline.base import StepFailed
from q3tui.pipeline.engine import Engine
from tests.test_rtl_flow import ScriptedLLM
from q3tui.eda.tools import preset_data
from tests.test_rtl_gen import calls, config_json, llm, spec_json, tools_dir, tools_yaml  # noqa: F401  (fixtures)

FAKE_TOOL = '''
import os, sys, time
import openpyxl
ws = openpyxl.load_workbook(sys.argv[1])[sys.argv[2]]
cfg = {str(r[0].value): r[1].value for r in ws.iter_rows(min_row=3, max_row=8)}
ports, regs = [], []
for r in ws.iter_rows(min_row=16):
    if r[0].value is not None:
        regs.append(r[0].value)
    elif r[2].value and r[2].value != "reserved":
        t = r[4].value
        ports += ([f"  input i_{regs[-1]}_{r[2].value}"] if t == "ro" else [f"  output o_{regs[-1]}_{r[2].value}"])
mod = cfg["Module Name"]
os.makedirs("RTL/models", exist_ok=True)
text = f"""`timescale 1ns/1ps
// Generator: someone
// Date and time: {time.time()}
module {mod} (
  input i_bus_clk,
  input i_bus_rstn,
  input [{int(cfg["Address Width"]) - 1}:0] i_paddr,
""" + ",\\n".join(ports + ["  output o_pready"]) + "\\n);\\n" + ("  m_vlsi_synch u_s ();\\n" if cfg["Asynchronous"] == 1 else "") + "endmodule\\n"
open(f"RTL/{mod}.sv", "w").write(text)
open("RTL/models/m_vlsi_synch.sv", "w").write("module m_vlsi_synch; endmodule\\n")
os.makedirs("Debug", exist_ok=True)
open("Debug/debug.json", "w").write("{}")
print("End of the CSR generation!")
'''



def _v2(src: str) -> str:
    """The fake tool as version 2.x of the generator writes: the synchronizer m_csr_sync, `input wire` / `output logic` ports with the new names, a
    block-comment header whose Author / Date lines change on every run."""
    for old, new in (("m_vlsi_synch", "m_csr_sync"),
                     ("// Generator: someone\n// Date and time: {time.time()}\n", "/*\n * Author       : someone\n * Date         : {time.time()}\n */\n"),
                     ("  input i_bus_clk,\n  input i_bus_rstn,\n  input [", "  input  wire i_clk_apb,\n  input  wire i_rst_n_apb,\n  input  wire ["),
                     ("i_paddr", "i_apb_paddr"), ('"  output o_pready"', '"  output logic o_apb_pready"')):
        assert old in src, old
        src = src.replace(old, new)
    return src


FAKE_TOOL_V2 = _v2(FAKE_TOOL)

SPEC_TEXT = "# blk\nThe block has a register map on an APB bus.\n| Offset | Register | Bits | Access | Reset |\n| 0x0 | ctrl | [0] en RW, [5:4] mode RW | | 0 |\n| 0x4 | stat | [0] busy RO, [1] done W1C | | 0 |\n"

MAP = {"has_registers": True, "module_name": "m_blk_csr", "addr_width": 16, "asynchronous": True, "req_ids": ["REQ-001"],
       "registers": [{"name": "ctrl", "offset": "0x0", "fields": [{"name": "en", "msb": 0, "lsb": 0, "type": "rw"},
                                                               {"name": "mode", "msb": 5, "lsb": 4, "type": "rwi", "reset": "0x1"}]},
                     {"name": "stat", "offset": 4, "fields": [{"name": "busy", "msb": 0, "lsb": 0, "type": "ro"}, {"name": "done", "msb": 1, "lsb": 1, "type": "w1c"}]}]}


def _map(**over):
    d = json.loads(json.dumps(MAP))
    d.update(over)
    return CsrMap.model_validate(d)


# -- the map: checks, workbook ----------------------------------------------------------------------------------------

def test_the_map_is_checked_because_the_tool_does_not():
    m = csr.complete(_map())
    assert csr.validate(m) == []
    assert [(f.name, f.msb, f.lsb) for f in m.registers[0].fields] == [("reserved", 31, 6), ("mode", 5, 4), ("reserved", 3, 1), ("en", 0, 0)]  # gaps filled
    bad = csr.validate(csr.complete(CsrMap.model_validate({"module_name": "x y", "addr_width": 8, "registers": [
        {"name": "a", "offset": 6, "fields": [{"name": "f", "msb": 3, "lsb": 0, "type": "rw"}, {"name": "g", "msb": 2, "lsb": 1, "type": "ro"}]},
        {"name": "A", "offset": 0x400, "fields": [{"name": "h", "msb": 40, "lsb": 36, "type": "rw"}, {"name": "i", "msb": 0, "lsb": 0, "type": "rw", "reset": 3}]},
        {"name": "b", "offset": 6, "fields": [{"name": "reserved", "msb": 0, "lsb": 0, "type": "rw"}]}]})))
    text = "\n".join(bad)
    for needle in ("not an identifier", "multiple of 4", "bit 1 is also in another field", "appears twice", "does not fit in 8 address bits", "share offset",
                   "not inside [31:0]", "reset value 0x3 does not fit", "`reserved` field is `ro`"):
        assert needle in text, needle
    assert csr.validate(_map(has_registers=False)) == [] and "supports 32 only" in "\n".join(csr.validate(csr.complete(_map(data_width=64))))
    assert csr.reset_value(csr.complete(_map()).registers[0]) == 0x10


def test_the_workbook_is_in_the_layout_the_tool_reads(tmp_path):
    wb = openpyxl.load_workbook(csr.write_workbook(csr.complete(_map()), tmp_path / "in.xlsx"))
    ws = wb["CSR"]
    rows = {r[0].row: [c.value for c in r[:7]] for r in ws.iter_rows() if any(c.value is not None for c in r[:7])}
    assert rows[1][0] == "Table - Configuration" and rows[3][:2] == ["Module Name", "m_blk_csr"] and rows[8][:2] == ["Asynchronous", 1]
    assert rows[7][:2] == ["Write Strobe", 4] and rows[6][:2] == ["Address Width", 16] and rows[5][:2] == ["Data Width", 32]
    assert rows[14][0] == "Table - Register Definition" and rows[15][:6] == ["REGISTER", "OFFSET", "BIT NAME", "FIELD WIDTH", "BIT TYPE", "RESET VALUE"]
    assert rows[16][:2] == ["ctrl", "0x0000"] and rows[16][5] == "0x10"  # the register's reset value follows its fields
    assert rows[17][2:6] == ["reserved", "[31:6]", "ro", "0x0"] and rows[18][2:6] == ["mode", "[5:4]", "rwi", "0x1"] and rows[20][2:5] == ["en", "[0]", "rw"]


def test_numbers_may_be_written_the_way_a_spec_does():
    m = CsrMap.model_validate({"module_name": "m", "registers": [{"name": "r", "offset": "16'h0010", "fields": [{"name": "f", "msb": "7", "lsb": "0", "type": "rw", "reset": "8'hA5"}]}]})
    assert (m.registers[0].offset, m.registers[0].fields[0].reset) == (16, 0xA5)


def test_ports_are_read_from_the_generated_text():
    sv = "module m_x (\n  input i_bus_clk,\n  input [15:0] i_paddr,\n  output o_pready,\n  input i_x_a,\n  output [1:0] o_x_b\n);\n  logic q;\nendmodule\n"
    assert csr.ports(sv) == [("input", "", "i_bus_clk"), ("input", "[15:0]", "i_paddr"), ("output", "", "o_pready"), ("input", "", "i_x_a"), ("output", "[1:0]", "o_x_b")]
    assert "register ports:\n input i_x_a\n output [1:0] o_x_b" in csr.port_summary(sv, "m_x").replace("  ", " ")
    assert csr.stable("// Generator: a\n// Date and time: 5\nmodule m;") == "\n\nmodule m;" and csr.mentions_registers("see the Register map") and not csr.mentions_registers("a fifo")
    # version 2.x of the tool: `wire` / `logic`, widths written with its localparams, new bus names, a block-comment header
    sv2 = ("module m_x #(\n  localparam logic [31:0] LPARA_ADDR_W = 32'd16,\n  localparam logic [31:0] LPARA_X_OFFSET = 32'h0004\n) (\n  // APB clock and reset\n"
           "  input  wire                     i_clk_apb,\n  input  wire  [LPARA_ADDR_W-1:0] i_apb_paddr,\n  output logic                    o_apb_pready,\n\n"
           "  // Register x\n  input  wire                     i_x_a,\n  output logic [1:0]              o_x_b\n);\nendmodule\n")
    assert csr.ports(sv2) == [("input", "", "i_clk_apb"), ("input", "[15:0]", "i_apb_paddr"), ("output", "", "o_apb_pready"), ("input", "", "i_x_a"), ("output", "[1:0]", "o_x_b")]
    assert "bus / clocks: i_clk_apb, i_apb_paddr, o_apb_pready\nregister ports:\n input i_x_a\n output [1:0] o_x_b" in csr.port_summary(sv2, "m_x").replace("  ", " ")
    assert csr.stable("/*\n * Author       : a\n * Date         : 2026-10-29\n */\nmodule m;") == "/*\n\n\n */\nmodule m;"


# -- the step ---------------------------------------------------------------------------------------------------------

@pytest.fixture
def csr_proj(tmp_path, monkeypatch):
    from q3tui.flows.spec.lib import artifacts

    monkeypatch.delenv("CSR_GEN_HOME", raising=False)  # (the tool is where the project says, not where this machine has one)
    tool = tmp_path / "gen"
    (tool / "models").mkdir(parents=True)
    (tool / "CSR_Generation.py").write_text(FAKE_TOOL)
    (tool / "models" / "m_vlsi_synch.sv").write_text("module m_vlsi_synch; endmodule\n")
    (tmp_path / "q3tui.yaml").write_text(f"pipeline: {{flow: rtl, step_options: {{csr: {{tool_dir: {tool}}}}}}}\n")
    (tmp_path / "config").mkdir()
    (tmp_path / "config" / "tools.json").write_text(json.dumps({"roles": {"csr_gen": preset_data()["roles"]["csr_gen"]}}))  # (the project's own tools file)
    (tmp_path / "spec").mkdir()
    (tmp_path / "spec" / "spec.md").write_text(SPEC_TEXT)
    project = Project.open(tmp_path)
    artifacts.write(project, "structured_spec.json", spec_json(), check=False)
    artifacts.write(project, "final_config.json", config_json(), check=False)
    fake = ScriptedLLM()
    fake.answers["csr_map"] = lambda s: _map()
    monkeypatch.setattr(runtime, "run_stage", fake)
    return tmp_path, Engine(project), fake


def run_csr(eng):
    anyio.run(lambda: eng.run_step(eng.by_name["csr"]))


def test_the_registers_come_from_the_tool_not_from_the_llm(csr_proj):
    root, eng, fake = csr_proj
    run_csr(eng)
    st = fake.calls[0]
    assert st.name == "csr_map" and {t.name for t in st.sdk_tools} == {"csr_check", "tools_available"} and "csr-register-map" in st.prompt and "apb-csr-tool" in st.prompt
    assert "ctrl | [0] en RW" in st.prompt and "leaf_a" in st.prompt and root / "src" / "rtl" not in [d for d in st.deny_dirs if d != root / "src" / "rtl"] and root / "src" / "rtl" in st.deny_dirs
    mp = json.loads((root / "src" / "rtl" / "csr" / "csr_map.json").read_text())
    assert mp["module_name"] == "m_blk_csr" and [f["name"] for f in mp["registers"][0]["fields"]][:2] == ["reserved", "mode"]  # (completed by code)
    text = (root / "src" / "rtl" / "m_blk_csr.sv").read_text()
    assert text.startswith("// Generated by the APB-CSR-Generator") and "// REQ-IDs: REQ-001" in text and "input i_stat_busy" in text and "u_s" in text
    assert (root / "src" / "rtl" / "m_vlsi_synch.sv").is_file() and (root / "src" / "rtl" / "csr" / "csr.xlsx").is_file()
    prov = csr.read_provided(root / "src" / "rtl")
    assert prov == {"m_blk_csr": {"kind": "csr", "parent": "top", "req_ids": ["REQ-001"], "instances": ["m_vlsi_synch"]},
                    "m_vlsi_synch": {"kind": "model", "parent": "m_blk_csr", "req_ids": [], "instances": []}}
    rep = json.loads((root / "src" / "rtl" / "csr" / "csr_report.json").read_text())
    assert rep["registers"] == ["ctrl", "stat"] and "i_stat_busy" in rep["ports"]
    # nothing changed: no LLM, and the tool's new time stamp is not a change of the module
    before = (root / "src" / "rtl" / "m_blk_csr.sv").read_text()
    time.sleep(0.01)
    eng.state.step_meta["csr"].pop("snapshot")  # (force a second look)
    fake.answers["csr_map"] = lambda s: _map()
    run_csr(eng)
    assert fake.names() == ["csr_map", "csr_map"]
    assert (root / "src" / "rtl" / "m_blk_csr.sv").read_text() == before
    n = len(fake.calls)
    run_csr(eng)
    assert len(fake.calls) == n  # the snapshot is the same: no LLM at all


def test_version_2_of_the_tool_replaces_the_module_and_its_synchronizer(csr_proj):
    """Version 2.0 of the generator renamed the synchronizer (m_csr_sync), the ports and the header: it is found, and a module made by 1.x is replaced,
    its synchronizer too; its own date is no change."""
    root, eng, fake = csr_proj
    run_csr(eng)  # (version 1.x)
    gen = root / "gen"
    (gen / "models" / "m_vlsi_synch.sv").unlink()
    (gen / "models" / "m_csr_sync.sv").write_text("module m_csr_sync; endmodule\n")
    (gen / "CSR_Generation.py").write_text(FAKE_TOOL_V2)
    assert eng.by_name["csr"].tool(eng) == gen and csr.missing(gen) == []
    eng.state.step_meta["csr"].pop("snapshot")
    run_csr(eng)
    rtl = root / "src" / "rtl"
    text = (rtl / "m_blk_csr.sv").read_text()
    assert "i_apb_paddr" in text and "m_csr_sync u_s" in text and (rtl / "m_csr_sync.sv").is_file() and not (rtl / "m_vlsi_synch.sv").exists()
    assert csr.read_provided(rtl)["m_blk_csr"]["instances"] == ["m_csr_sync"] and "m_vlsi_synch" not in csr.read_provided(rtl)
    assert "i_apb_paddr" in json.loads((rtl / "csr" / "csr_report.json").read_text())["ports"]
    time.sleep(0.01)
    eng.state.step_meta["csr"].pop("snapshot")
    run_csr(eng)
    assert (rtl / "m_blk_csr.sv").read_text() == text  # (a new Date line is not a change)


def test_you_own_the_map_after_the_first_draft(csr_proj):
    root, eng, fake = csr_proj
    run_csr(eng)
    f = root / "src" / "rtl" / "csr" / "csr_map.json"
    mp = json.loads(f.read_text())
    mp["registers"][1]["offset"] = 8
    f.write_text(json.dumps(mp))
    assert eng.evaluate(eng.by_name["csr"]).status == "stale"
    run_csr(eng)
    assert len(fake.calls) == 1 and json.loads((root / "src" / "rtl" / "csr" / "csr_report.json").read_text())["registers"] == ["ctrl", "stat"]
    wb = openpyxl.load_workbook(root / "src" / "rtl" / "csr" / "csr.xlsx")["CSR"]
    assert any(r[1].value == "0x0008" for r in wb.iter_rows(min_row=16))
    mp["registers"][1]["offset"] = 0  # your edit collides with `ctrl`: the tool would silently build broken RTL
    f.write_text(json.dumps(mp))
    with pytest.raises(StepFailed, match="share offset 0x0"):
        run_csr(eng)


def test_a_map_with_problems_goes_back_to_the_llm_and_the_spec_changing_updates_it(csr_proj):
    root, eng, fake = csr_proj
    bad = _map(registers=[{"name": "a", "offset": 0, "fields": [{"name": "x", "msb": 3, "lsb": 0, "type": "rw"}, {"name": "y", "msb": 2, "lsb": 0, "type": "rw"}]}])
    answers = iter([bad])
    fake.answers["csr_map"] = lambda s: next(answers)
    fake.answers["csr_map_fix"] = lambda s: _map()
    run_csr(eng)
    assert fake.names() == ["csr_map", "csr_map_fix"] and "also in another field" in fake.calls[1].prompt
    (root / "spec" / "spec.md").write_text(SPEC_TEXT + "| 0x8 | irq | [0] pending W1C | | 0 |\n")
    fake.answers["csr_map"] = lambda s: _map(summary="irq added")
    run_csr(eng)
    last = fake.calls[-1].prompt
    assert fake.calls[-1].name == "csr_map" and "the spec changed since: update only what the changes require" in last and '"module_name": "m_blk_csr"' in last
    always_bad = lambda s: bad  # noqa: E731
    fake.answers["csr_map"], fake.answers["csr_map_fix"] = always_bad, always_bad
    with pytest.raises(StepFailed, match="problems after 3 tries"):
        anyio.run(lambda: eng.run_step(eng.by_name["csr"], regenerate=True))


def test_no_register_no_llm_no_tool_and_a_missing_tool_is_said(csr_proj):
    root, eng, fake = csr_proj
    (root / "spec" / "spec.md").write_text("# fifo\nA plain fifo with a push and a pop port.\n")
    run_csr(eng)
    assert fake.calls == [] and json.loads((root / "src" / "rtl" / "csr" / "csr_map.json").read_text())["has_registers"] is False and csr.read_provided(root / "src" / "rtl") == {}
    assert eng.by_name["csr"].missing_input(eng) is None  # (nothing to generate: the tool is not needed)
    (root / "spec" / "spec.md").write_text(SPEC_TEXT)
    run_csr(eng)
    (root / "q3tui.yaml").write_text("pipeline: {flow: rtl, step_options: {csr: {tool_dir: /nonexistent}}}\n")
    eng2 = Engine(Project.open(root))
    msg = eng2.by_name["csr"].missing_input(eng2)
    assert "csr_tool_install" in msg and "not installed" in msg
    eng2.by_name["csr"].options["mode"] = "off"
    assert eng2.by_name["csr"].missing_input(eng2) is None
    anyio.run(lambda: eng2.run_step(eng2.by_name["csr"]))
    assert csr.read_provided(root / "src" / "rtl") == {} and not (root / "src" / "rtl" / "m_blk_csr.sv").exists()  # off: what an earlier run made goes (backed up)
    assert list((root / ".q3tui" / "bkp").glob("*/src/rtl/m_blk_csr.sv"))


def test_the_csr_check_tool_validates_and_lists_the_ports(csr_proj):
    from q3tui.flows.rtl.lib.tools import csr_tools

    root, eng, fake = csr_proj
    step = eng.by_name["csr"]
    tool = {t.name: t for t in csr_tools(step.tools(eng), step.tool(eng), root / ".q3tui" / "w")}["csr_check"]

    def call(m):
        return anyio.run(lambda: tool.handler({"map": m}))["content"][0]["text"]

    assert "share offset" in call({**MAP, "registers": [MAP["registers"][0], {**MAP["registers"][1], "offset": 0}]})
    out = call(MAP)
    assert out.startswith("the map is valid and generates:") and "i_stat_busy" in out and "module m_blk_csr" in out
    assert "not in the expected structure" in call({"registers": "nope"})


def test_a_changed_generated_module_is_kept_and_its_patch_is_applied_to_the_next_one(csr_proj):
    """The rtl step changes the generated module (lint fixes, change requests): the csr step keeps that file while the tool makes the same module,
    and when the map changes the new module gets the changes again from their patch (src/rtl/csr/patches/), by code."""
    from q3tui.flows.rtl.lib import patches

    root, eng, fake = csr_proj
    run_csr(eng)
    rtl = root / "src" / "rtl"
    sv = rtl / "m_blk_csr.sv"
    assert sv.read_text().endswith("endmodule\n")                                   # (a newline at the end: no EOFNEWLINE)
    assert patches.base(eng.project, "m_blk_csr") == sv.read_text()                  # what the tool made: the base of the patch
    fixed = sv.read_text().replace("  input i_bus_rstn,\n", "  /* verilator lint_off UNUSEDSIGNAL */\n  input i_bus_rstn,\n")
    sv.write_text(fixed)
    patches.record(eng.project, rtl, "m_blk_csr", ["lint fix: i_bus_rstn is unused"])  # (as the rtl step does after a change)
    assert "+  /* verilator lint_off UNUSEDSIGNAL */" in patches.patch_path(rtl, "m_blk_csr").read_text()
    assert "lint fix: i_bus_rstn is unused" in patches.notes_path(rtl, "m_blk_csr").read_text()
    run_csr(eng)                                                                     # same map: the tool makes the same module
    assert sv.read_text() == fixed
    seen = []
    eng.bus.subscribe(lambda ev: seen.append(ev))
    mp = rtl / "csr" / "csr_map.json"
    data = json.loads(mp.read_text())
    data["registers"].append({"name": "cfg", "offset": "0x8", "fields": [{"name": "k", "msb": 0, "lsb": 0, "type": "rw"}]})
    mp.write_text(json.dumps(data))
    run_csr(eng)                                                                     # your new map: a new module, with the change applied again
    assert "lint_off UNUSEDSIGNAL" in sv.read_text() and "o_cfg_k" in sv.read_text()
    assert any("applied to it" in str(getattr(ev, "data", ev)) for ev in seen) and not eng.state.feedback.get("rtl")
    assert "o_cfg_k" in patches.base(eng.project, "m_blk_csr") and "lint_off" not in patches.base(eng.project, "m_blk_csr")
    assert "+  /* verilator lint_off UNUSEDSIGNAL */" in patches.patch_path(rtl, "m_blk_csr").read_text()   # (the patch, now against the new module)
    assert "generated again" in patches.notes_path(rtl, "m_blk_csr").read_text()


def test_a_change_that_no_longer_fits_the_new_module_goes_to_the_rtl_step(csr_proj):
    from q3tui.flows.rtl.lib import patches

    root, eng, fake = csr_proj
    run_csr(eng)
    rtl = root / "src" / "rtl"
    sv = rtl / "m_blk_csr.sv"
    gen = sv.read_text()
    patches.save_base(eng.project, "m_blk_csr", gen.replace("endmodule\n", "  // OLD LINE A\n  // OLD LINE B\n  // OLD LINE C\nendmodule\n"))
    sv.write_text(gen.replace("endmodule\n", "  // OLD LINE A\n  // NEW LINE\n  // OLD LINE C\nendmodule\n"))  # (a change to lines the next module has no more)
    mp = rtl / "csr" / "csr_map.json"
    data = json.loads(mp.read_text())
    data["registers"].append({"name": "cfg", "offset": "0x8", "fields": [{"name": "k", "msb": 0, "lsb": 0, "type": "rw"}]})
    mp.write_text(json.dumps(data))
    run_csr(eng)
    assert "NEW LINE" not in sv.read_text() and "o_cfg_k" in sv.read_text()
    (req,) = eng.state.feedback["rtl"]
    assert req.startswith("[rtl:m_blk_csr]") and "no longer fit" in req and "+  // NEW LINE" in req and "m_blk_csr.md" in req


def test_a_patch_applies_where_its_lines_moved_to_and_the_tool_s_date_is_no_change():
    from q3tui.flows.rtl.lib import patches

    base = "// Date and time: 2026-01-01 10:00\nmodule m (\n  input a,\n  input b\n);\n  assign x = a;\nendmodule\n"
    fixed = base.replace("  assign x = a;\n", "  assign x = a & b;\n")
    p = patches.diff(base, fixed, "m")
    assert patches.diff(base, base.replace("2026-01-01 10:00", "2026-02-02 11:00"), "m") == ""   # the date alone is no change
    new = "// header line\n" + base.replace("2026-01-01 10:00", "2026-10-28 21:00").replace("  input b\n", "  input b,\n  input c\n")
    out, failed = patches.apply(new, p)                                             # (a line more above it, a port more next to it, another date)
    assert not failed and "  assign x = a & b;\n" in out and "2026-10-28 21:00" in out and "input c" in out and out.startswith("// header line\n")
    out, failed = patches.apply(new.replace("  assign x = a;\n", "  assign x = 0;\n"), p)
    assert len(failed) == 1 and "+  assign x = a & b;" in failed[0] and "assign x = 0;" in out


# -- the rtl step with a provided module ---------------------------------------------------------------------------------

GEN = "module m_blk_csr (input i_bus_clk);\n  m_vlsi_synch u_s ();\n  // LINT_WARN\nendmodule\n"


def _provided_project(tmp_path, tools_dir, opts: str = "") -> Engine:
    from q3tui.flows.spec.lib import artifacts

    (tmp_path / "q3tui.yaml").write_text(f"pipeline: {{flow: rtl, step_options: {{rtl: {{rules: [rtl_rule.md]{opts}}}}}}}\n" + tools_yaml(tools_dir))
    (tmp_path / "spec").mkdir()
    (tmp_path / "spec" / "spec.md").write_text("# demo spec\n")
    project = Project.open(tmp_path)
    artifacts.write(project, "structured_spec.json", spec_json(), check=False)
    artifacts.write(project, "final_config.json", config_json(), check=False)
    rtl = tmp_path / "src" / "rtl"
    rtl.mkdir(parents=True)
    (rtl / "m_blk_csr.sv").write_text(GEN)
    (rtl / "m_vlsi_synch.sv").write_text("module m_vlsi_synch; endmodule\n")
    (rtl / "csr").mkdir()
    (rtl / "csr" / "provided.json").write_text(json.dumps({"m_blk_csr": {"kind": "csr", "parent": "top", "req_ids": ["REQ-001"], "instances": ["m_vlsi_synch"]},
                                                          "m_vlsi_synch": {"kind": "model", "parent": "m_blk_csr", "req_ids": [], "instances": []}}))
    return Engine(project)


def test_the_rtl_step_fixes_the_lint_of_the_provided_module_in_place(tmp_path, tools_dir, llm, monkeypatch):
    """Generated code is never written from the spec nor rule-checked, but its lint findings are fixed in place (name and ports kept)."""
    eng = _provided_project(tmp_path, tools_dir)
    rtl = tmp_path / "src" / "rtl"
    eng.state.step("csr").outputs["src/rtl/m_blk_csr.sv"] = "old"  # (as the csr step recorded it)

    async def spy(stage, cfg, emit):
        if stage.name == "rtl_m_blk_csr":
            stage.write_files[0].write_text(stage.write_files[0].read_text().replace("  // LINT_WARN\n", ""))
            from q3tui.flows.rtl.rtl.gen import VModuleResult
            from q3tui.llm.runtime import StageResult
            return StageResult("", VModuleResult(summary="fixed"), 0.01, 2, "s")
        return await llm(stage, cfg, emit)

    monkeypatch.setattr(runtime, "run_stage", spy)
    anyio.run(lambda: eng.run_step(eng.by_name["rtl"]))
    text = (rtl / "m_blk_csr.sv").read_text()
    assert "LINT_WARN" not in text and text.startswith("module m_blk_csr (input i_bus_clk);")
    assert not [c for c in llm.calls if c.name == "rtl_m_vlsi_synch"]               # clean: nothing to fix
    assert eng.state.step("csr").outputs["src/rtl/m_blk_csr.sv"] != "old"           # the csr step's output, fixed here: not "edited"
    assert eng.state.steps["rtl"].status == "done"


def test_a_change_request_for_a_generated_module_is_applied_in_place_and_kept_as_a_patch(tmp_path, tools_dir, llm, monkeypatch):
    """A simulation's RTL fix for the generated module ([rtl:m_blk_csr] …): a small session changes it in place; the patch and its why are recorded."""
    from q3tui.flows.rtl.lib import patches

    eng = _provided_project(tmp_path, tools_dir)
    rtl = tmp_path / "src" / "rtl"
    patches.save_base(eng.project, "m_blk_csr", GEN)
    eng.request_change("rtl", "[rtl:m_blk_csr] commit a write at the edge ending the ACCESS cycle with PREADY = 1")
    prompts = []

    async def spy(stage, cfg, emit):
        if stage.name == "rtl_m_blk_csr":
            prompts.append(stage.prompt)
            f = stage.write_files[0]
            f.write_text(f.read_text().replace("  // LINT_WARN\n", "") .replace("endmodule", "  // commit at PREADY\nendmodule"))
            from q3tui.flows.rtl.rtl.gen import VModuleResult
            from q3tui.llm.runtime import StageResult
            return StageResult("", VModuleResult(summary="changed"), 0.01, 2, "s")
        return await llm(stage, cfg, emit)

    monkeypatch.setattr(runtime, "run_stage", spy)
    anyio.run(lambda: eng.run_step(eng.by_name["rtl"]))
    assert "commit a write at the edge ending the ACCESS cycle" in prompts[0] and "GIỮ NGUYÊN tên module" in prompts[0]
    assert "// commit at PREADY" in (rtl / "m_blk_csr.sv").read_text()
    assert "+  // commit at PREADY" in patches.patch_path(rtl, "m_blk_csr").read_text()
    assert "change request: commit a write at the edge ending the ACCESS cycle" in patches.notes_path(rtl, "m_blk_csr").read_text()


def test_with_keep_generated_a_change_request_for_a_generated_module_is_only_reported(tmp_path, tools_dir, llm):
    eng = _provided_project(tmp_path, tools_dir, ", keep_generated: true")
    eng.request_change("rtl", "[rtl:m_blk_csr] commit a write at the edge ending the ACCESS cycle with PREADY = 1")
    seen = []
    eng.bus.subscribe(lambda ev: seen.append(ev))
    anyio.run(lambda: eng.run_step(eng.by_name["rtl"]))
    assert (tmp_path / "src" / "rtl" / "m_blk_csr.sv").read_text() == GEN and not [c for c in llm.calls if "m_blk_csr" in c.name]
    assert any("not applied" in str(getattr(ev, "data", ev)) for ev in seen)


def test_the_fix_prompt_of_a_generated_module_keeps_its_interface():
    from q3tui.flows.rtl.rtl import task

    p = task.fix_prompt(module="m_blk_csr", path="src/rtl/m_blk_csr.sv", problems="- lint warning", code="module m;", reqs=[], generated=True)
    assert "GIỮ NGUYÊN tên module" in p and "lint_off UNUSEDSIGNAL" in p and "REQ-ID" not in p


def test_the_rtl_step_instantiates_the_provided_module_and_never_writes_it(tmp_path, tools_dir, llm):
    """Option keep_generated: the module stays as the generator made it (its lint findings are reported only)."""
    eng = _provided_project(tmp_path, tools_dir, ", keep_generated: true")
    rtl, gen = tmp_path / "src" / "rtl", GEN
    anyio.run(lambda: eng.run_step(eng.by_name["rtl"]))
    assert not [c for c in llm.calls if "m_blk_csr" in c.name or "m_vlsi_synch" in c.name]  # no stage writes or fixes them
    assert (rtl / "m_blk_csr.sv").read_text() == gen
    assert (rtl / "filelist.f").read_text().split()[-3:] == ["src/rtl/m_vlsi_synch.sv", "src/rtl/m_blk_csr.sv", "src/rtl/top.sv"] or "src/rtl/m_blk_csr.sv" in (rtl / "filelist.f").read_text().split()
    step = eng.by_name["rtl"]
    assert not [f for f in step.outputs(eng) if f.stem in ("m_blk_csr", "m_vlsi_synch")]  # the csr step's files are not this step's outputs
    assert tmp_path / "src" / "rtl" / "csr" / "provided.json" in step.inputs(eng)
    # the flow's own steps leave generated code alone, and the smoke test compiles the whole design
    fmt = eng.by_name["format"]
    assert sorted(f.stem for f in fmt.rtl_files(eng)) == ["leaf_a", "leaf_b", "top"]
    assert {"m_blk_csr", "m_vlsi_synch"} <= {f.stem for f in eng.by_name["smoke"].design_files(eng)}
    assert (tmp_path / "src" / "rtl") in [d for d in eng.by_name["smoke"]._deny(eng)] or (tmp_path / "src" / "rtl" / "csr") in eng.by_name["smoke"]._deny(eng)
    # a lint finding in generated code is reported, never sent to an LLM to rewrite
    assert not any(c.name.startswith("rtl_fix") and "csr" in c.prompt for c in llm.calls)


# -- the real generator ---------------------------------------------------------------------------------------------------

REAL = csr.find_tool(Path("/nonexistent"))  # ($CSR_GEN_HOME or ~/.q3tui/tools/APB-CSR-Generator, at collection time)


@pytest.mark.skipif(REAL is None, reason="the APB-CSR-Generator is not installed ($CSR_GEN_HOME or ~/.q3tui/tools/APB-CSR-Generator)")
def test_the_real_generator(tmp_path):
    from q3tui.core.config import ToolRole
    from q3tui.eda.tools import ToolSet, ToolEnv

    m = csr.complete(_map())
    assert csr.validate(m) == []
    work = tmp_path / "w"
    wb = csr.write_workbook(m, work / "in.xlsx")
    flow_tools = preset_data()["roles"]["csr_gen"]
    ts = ToolSet({"csr_gen": ToolRole.model_validate(flow_tools)}, ToolEnv(None, (), None), 120)
    res = csr.run_tool(ts, REAL, wb, work)
    assert res.ok, res.summary
    sv = (work / "RTL" / "m_blk_csr.sv").read_text()
    names = [n for _, _, n in csr.ports(sv)]
    bus = {"i_clk_apb", "o_apb_pready", "i_clk_reg"} if csr.sync_model(work / "RTL") == "m_csr_sync" else {"i_bus_clk", "o_pready", "i_reg_clk"}  # (v2.x, v1.x)
    assert bus | {"o_ctrl_en", "o_ctrl_mode", "i_hw_wdata_ctrl_mode", "i_hw_we_ctrl_mode", "i_stat_busy", "o_stat_done", "i_hw_we_stat_done"} <= set(names)
    assert not any("reserved" in n for n in names) and csr.sync_model(work / "RTL") == csr.sync_model(REAL)
    assert set(csr.port_summary(sv, "m_blk_csr").split("bus / clocks: ")[1].split("\n")[0].split(", ")) >= bus  # (the bus ports are told apart)


def test_the_generator_is_found_through_the_csr_gen_role_env(tmp_path, monkeypatch):
    """CSR_GEN_HOME in the `csr_gen` role's env (tools.json) is enough: it need not be exported in the shell."""
    monkeypatch.delenv("CSR_GEN_HOME", raising=False)
    monkeypatch.setenv("Q3TUI_HOME", str(tmp_path / "home"))
    tool = tmp_path / "gen"
    (tool / "models").mkdir(parents=True)
    (tool / csr.SCRIPT).write_text("")
    (tool / "models" / f"{csr.SYNC_MODELS[0]}.sv").write_text("")
    assert csr.find_tool(tmp_path) is None
    assert csr.find_tool(tmp_path, None, str(tool)) == tool
    assert csr.find_tool(tmp_path, str(tmp_path / "nope"), str(tool)) == tool  # (an option that points nowhere is skipped)
